import duckdb
from openpyxl import load_workbook
from time import time

BRONZE_PATH = "data/bronze/tabela5938.xlsx"
SILVER_CREATE_PATH = "sql/ddl/create_silver.sql"
SILVER_TRUNC_PATH = "sql/ddl/truncate_silver.sql"
SILVER_DML_PATH = "sql/dml/silver_dml.sql"

def create_silver():
    
    START_TIME = time()
    
    conn = duckdb.connect(":memory:")
    conn.execute("INSTALL excel; LOAD excel;")
    
    ##### Bronze #####
                                    # "_1", "_2", ... "_20"
    old_anos_cols = ["Ano", "C2"] + [f"\"_{x}\"" for x in range(1,21)]
    new_anos_cols = [f"ano_{y}" for y in range(2002, 2024)]
    
    assert len(old_anos_cols) == len(new_anos_cols)
    
    print(f"[INFO] Reading raw data from {BRONZE_PATH}...")
    conn.execute("CREATE SCHEMA bronze")
    
    wb = load_workbook(filename=BRONZE_PATH, read_only=True, keep_links=False)
    
    try:
        for table, sheet in {
            "pib": wb.sheetnames[0],
            "impostos": wb.sheetnames[1],
            "valor_adicionado": wb.sheetnames[2]
        }.items():
            bronze_table = f"bronze.{table}"
            
            conn.execute(f"""
                CREATE TABLE {bronze_table} AS
                SELECT
                    "Unidade da Federação e Município" AS localidade,
                    {", ".join([
                        f"{old_anos_cols[i]} AS {new_anos_cols[i]}"
                        for i in range(len(old_anos_cols))
                    ])}
                FROM read_xlsx(
                    '{BRONZE_PATH}',
                    sheet='{sheet}',
                    range='A3:W5601',
                    header=true,
                    all_varchar=true
                );
            """)
    finally:
        wb.close()
    
    ##### Silver #####
    
    print("[INFO] Creating tables for silver layer...")
    
    for path in [SILVER_CREATE_PATH, SILVER_DML_PATH]:
        with open(path, mode="r") as f:
            conn.execute(f.read())
    
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS silver.municipios AS (
            WITH municipio AS (
                SELECT
                    SUBSTRING(localidade FROM 1 FOR LENGTH(localidade) - 5) AS nome_municipio,
                    RIGHT(localidade, 3)[1:2] AS sigla_uf
                FROM bronze_pib
                WHERE localidade LIKE '%)'
            )
            
            SELECT
                ROW_NUMBER() OVER(
                    ORDER BY m.nome_municipio, u.uf_id
                ) AS municipio_id,
                m.nome_municipio,
                u.uf_id
            FROM municipio m
            JOIN silver.ufs u
                ON u.sigla_uf = m.sigla_uf
            GROUP BY m.nome_municipio, u.uf_id
        )
    """)
    
    print("[INFO] Creating factual tables for silver layer...")
    
    for s in sheets.keys():
        bronze_table = f"bronze_{s}"
        silver_table = f"silver.fact_{s}"
        valor_col = f"valor_{s}"
        
        conn.execute(f"""
            CREATE OR REPLACE TABLE {silver_table} AS (
                WITH bronze_unpivot AS (
                    SELECT *
                    FROM {bronze_table}
                    UNPIVOT (valor FOR coluna_ano IN ({", ".join(new_anos_cols)}))
                    WHERE localidade LIKE '%)'
                ),
                
                municipio_uf AS (
                    SELECT
                        m.municipio_id,
                        m.nome_municipio,
                        u.sigla_uf
                    FROM silver.municipios m
                    JOIN silver.ufs u
                        ON u.uf_id = m.uf_id
                )
                
                SELECT
                    m.municipio_id,
                    a.ano_id,
                    CASE
                        WHEN b.valor = '...' THEN NULL
                        ELSE CAST(b.valor AS DOUBLE)
                    END AS {valor_col}
                FROM bronze_unpivot b
                JOIN municipio_uf m
                    ON m.nome_municipio = SUBSTRING(b.localidade FROM 1 FOR LENGTH(b.localidade) - 5)
                    AND m.sigla_uf = RIGHT(b.localidade, 3)[1:2]
                JOIN silver.anos a
                    ON a.ano = CAST(REPLACE(b.coluna_ano, 'ano_', '') AS INTEGER)
            )
        """)
    
    END_TIME = time()
    
    print(f"[SUCCESS] Bronze -> Silver process finished successfully in {END_TIME - START_TIME:.2f} seconds!")